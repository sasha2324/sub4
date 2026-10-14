# ============================================
# STRAVA — REPORT EXCEL
# ============================================
# Legge strava_corse.csv e genera strava_report.xlsx
# con 4 fogli: dashboard, elenco corse, riepilogo
# mensile e record personali.
#
# Prima di usarlo:  pip3 install openpyxl
# Esegui:           python3 strava_report.py
# ============================================

import csv
import os
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FILE_CSV = "strava_corse.csv"
FILE_XLSX = "strava_report.xlsx"

# Il report copre solo la stagione in corso. Con --tutto si torna
# a includere anche l'archivio 2025.
DAL = "2026-01-01"

# Colori: l'arancione è quello di Strava
ARANCIONE = "FC4C02"
GRIGIO     = "F2F2F2"
GRIGIO_SC  = "D9D9D9"
TESTO      = "1A1A1A"

MESI_IT = ["gen", "feb", "mar", "apr", "mag", "giu",
           "lug", "ago", "set", "ott", "nov", "dic"]

# ============================================
# LETTURA E CALCOLI
# ============================================

def carica_corse(nome_file, dal=None):
    """Legge il CSV prodotto dal tracker o dall'importatore."""
    if not os.path.exists(nome_file):
        print(f"✗ File '{nome_file}' non trovato.")
        print("  Esegui prima strava_import_export.py con l'archivio Strava.")
        exit(1)

    corse = []
    with open(nome_file, "r", encoding="utf-8") as f:
        for riga in csv.DictReader(f):
            if dal and riga["data"] < dal:
                continue
            try:
                corse.append({
                    "data":       riga["data"],
                    "data_dt":    datetime.strptime(riga["data"], "%Y-%m-%d"),
                    "nome":       riga["nome"],
                    "km":         float(riga["km"]),
                    "minuti":     int(riga["minuti"]),
                    "passo":      riga["passo"],
                    "dislivello": int(riga.get("dislivello", 0) or 0),
                    "bpm":        riga.get("bpm_media", "N/D"),
                })
            except (ValueError, KeyError):
                continue

    if not corse:
        print("✗ Nessuna corsa valida nel CSV.")
        exit(1)

    corse.sort(key=lambda c: c["data_dt"], reverse=True)
    print(f"✓ Lette {len(corse)} corse da '{nome_file}'.")
    return corse

def passo_da_secondi(secondi):
    return f"{int(secondi) // 60}:{int(secondi) % 60:02d}"

def durata_leggibile(minuti):
    """34 → '34min'   |   62 → '1h02'"""
    if minuti < 60:
        return f"{minuti}min"
    return f"{minuti // 60}h{minuti % 60:02d}"

def etichetta_mese(chiave):
    """'2026-08' → 'ago 26'"""
    anno, mese = chiave.split("-")
    return f"{MESI_IT[int(mese) - 1]} {anno[2:]}"

def riepilogo_mensile(corse):
    """Aggrega per mese, dal più recente al più vecchio."""
    mesi = defaultdict(lambda: {"km": 0.0, "uscite": 0, "minuti": 0})
    for c in corse:
        m = mesi[c["data"][:7]]
        m["km"]     += c["km"]
        m["uscite"] += 1
        m["minuti"] += c["minuti"]
    return dict(sorted(mesi.items(), reverse=True))

def record_personali(corse):
    """Miglior passo per ciascuna distanza classica."""
    distanze = {
        "5 km":        (4.5,  5.5),
        "10 km":       (9.0,  11.0),
        "Mezza (21k)": (19.0, 23.0),
        "Maratona":    (40.0, 45.0),
    }
    records = {}
    for etichetta, (minimo, massimo) in distanze.items():
        candidate = [c for c in corse if minimo <= c["km"] <= massimo]
        if candidate:
            migliore = min(candidate, key=lambda c: c["minuti"] / c["km"])
            records[etichetta] = migliore
    return records

# ============================================
# STILI
# ============================================

def stile_titolo(cella, dimensione=18):
    cella.font = Font(bold=True, size=dimensione, color=TESTO)

def stile_intestazione(cella):
    cella.font = Font(bold=True, size=11, color="FFFFFF")
    cella.fill = PatternFill("solid", fgColor=ARANCIONE)
    cella.alignment = Alignment(horizontal="center", vertical="center")

def bordo_sottile():
    lato = Side(style="thin", color=GRIGIO_SC)
    return Border(left=lato, right=lato, top=lato, bottom=lato)

def larghezze(foglio, misure):
    for i, larghezza in enumerate(misure, start=1):
        foglio.column_dimensions[get_column_letter(i)].width = larghezza

def scrivi_tabella(foglio, riga_inizio, intestazioni, righe):
    """Scrive un'intestazione colorata e le righe sotto, a righe alternate."""
    for colonna, testo in enumerate(intestazioni, start=1):
        stile_intestazione(foglio.cell(row=riga_inizio, column=colonna, value=testo))

    bordo = bordo_sottile()
    for i, valori in enumerate(righe):
        r = riga_inizio + 1 + i
        for colonna, valore in enumerate(valori, start=1):
            cella = foglio.cell(row=r, column=colonna, value=valore)
            cella.border = bordo
            if i % 2 == 1:
                cella.fill = PatternFill("solid", fgColor=GRIGIO)
    return riga_inizio + len(righe) + 1

# ============================================
# I QUATTRO FOGLI
# ============================================

def foglio_dashboard(wb, corse, mesi):
    f = wb.active
    f.title = "📊 Dashboard"
    larghezze(f, [16, 16, 16, 16, 16, 16])

    km_totali  = sum(c["km"] for c in corse)
    min_totali = sum(c["minuti"] for c in corse)
    passo_medio = min_totali * 60 / km_totali if km_totali else 0

    stile_titolo(f.cell(row=1, column=1, value="🏃 STRAVA RUNNING REPORT"))
    f.cell(row=2, column=1,
           value=f"Generato il {datetime.now():%d/%m/%Y}  •  {len(corse)} corse  •  "
                 f"dal {corse[-1]['data']} al {corse[0]['data']}").font = Font(
        size=10, color="666666")

    # Riga dei numeri chiave
    intestazioni = ["Corse totali", "Km totali", "Tempo totale",
                    "Passo medio", "Uscita più lunga", "Media/uscita"]
    valori = [
        len(corse),
        round(km_totali, 1),
        f"{min_totali // 60}h {min_totali % 60}m",
        passo_da_secondi(passo_medio),
        round(max(c["km"] for c in corse), 1),
        round(km_totali / len(corse), 1),
    ]
    unita = ["uscite", "km", "", "min/km", "km", "km"]

    for colonna, (testo, valore, misura) in enumerate(
            zip(intestazioni, valori, unita), start=1):
        stile_intestazione(f.cell(row=4, column=colonna, value=testo))

        cella = f.cell(row=5, column=colonna, value=valore)
        cella.font = Font(bold=True, size=20, color=ARANCIONE)
        cella.alignment = Alignment(horizontal="center")

        misura_cella = f.cell(row=6, column=colonna, value=misura)
        misura_cella.font = Font(size=9, color="888888")
        misura_cella.alignment = Alignment(horizontal="center")

    # Dati dei grafici, in fondo al foglio (li leggono i grafici qui sotto)
    riga_dati = 30
    f.cell(row=riga_dati, column=1, value="Mese")
    f.cell(row=riga_dati, column=2, value="Km")
    f.cell(row=riga_dati, column=3, value="Uscite")
    f.cell(row=riga_dati, column=4, value="Passo (sec/km)")

    cronologico = sorted(mesi.items())
    for i, (chiave, dati) in enumerate(cronologico, start=1):
        passo_sec = dati["minuti"] * 60 / dati["km"] if dati["km"] else 0
        f.cell(row=riga_dati + i, column=1, value=etichetta_mese(chiave))
        f.cell(row=riga_dati + i, column=2, value=round(dati["km"], 1))
        f.cell(row=riga_dati + i, column=3, value=dati["uscite"])
        f.cell(row=riga_dati + i, column=4, value=round(passo_sec))

    ultima = riga_dati + len(cronologico)
    etichette = Reference(f, min_col=1, min_row=riga_dati + 1, max_row=ultima)

    grafico_km = BarChart()
    grafico_km.title = "Km per mese"
    grafico_km.height, grafico_km.width = 8, 16
    grafico_km.add_data(Reference(f, min_col=2, min_row=riga_dati, max_row=ultima),
                        titles_from_data=True)
    grafico_km.set_categories(etichette)
    grafico_km.y_axis.title = "km"
    f.add_chart(grafico_km, "A8")

    grafico_passo = LineChart()
    grafico_passo.title = "Passo medio mensile (sec/km) — più basso = più veloce"
    grafico_passo.height, grafico_passo.width = 8, 16
    grafico_passo.add_data(Reference(f, min_col=4, min_row=riga_dati, max_row=ultima),
                           titles_from_data=True)
    grafico_passo.set_categories(etichette)
    grafico_passo.y_axis.title = "secondi al km"
    f.add_chart(grafico_passo, "A25")

def foglio_corse(wb, corse):
    f = wb.create_sheet("📋 Tutte le corse")
    larghezze(f, [12, 30, 8, 10, 10, 14, 8])

    righe = [[
        c["data"], c["nome"], c["km"], durata_leggibile(c["minuti"]),
        c["passo"], c["dislivello"],
        int(c["bpm"]) if str(c["bpm"]).isdigit() else "—",
    ] for c in corse]

    scrivi_tabella(f, 1,
                   ["Data", "Nome", "Km", "Durata", "Passo",
                    "Dislivello (m)", "BPM"],
                   righe)
    f.freeze_panes = "A2"

def foglio_mensile(wb, corse, mesi):
    f = wb.create_sheet("📅 Per mese")
    larghezze(f, [12, 10, 12, 12, 14, 12])

    righe = []
    for chiave, dati in mesi.items():
        passo = dati["minuti"] * 60 / dati["km"] if dati["km"] else 0
        righe.append([
            etichetta_mese(chiave), dati["uscite"], round(dati["km"], 1),
            round(dati["km"] / dati["uscite"], 1),
            durata_leggibile(dati["minuti"]), passo_da_secondi(passo),
        ])

    km_totali  = sum(c["km"] for c in corse)
    min_totali = sum(c["minuti"] for c in corse)
    righe.append(["TOTALE", len(corse), round(km_totali, 1), "—",
                  durata_leggibile(min_totali), "—"])

    ultima = scrivi_tabella(f, 1, ["Mese", "Uscite", "Km totali", "Km media",
                                   "Tempo totale", "Passo medio"], righe)

    for colonna in range(1, 7):
        f.cell(row=ultima - 1, column=colonna).font = Font(bold=True)

def foglio_record(wb, corse):
    f = wb.create_sheet("🏆 Record personali")
    larghezze(f, [26, 16, 16, 12, 14])

    stile_titolo(f.cell(row=1, column=1, value="🏆 I TUOI RECORD PERSONALI"), 14)

    records = record_personali(corse)
    righe = []
    for etichetta in ["5 km", "10 km", "Mezza (21k)", "Maratona"]:
        c = records.get(etichetta)
        if c:
            passo = c["minuti"] * 60 / c["km"]
            righe.append([etichetta, round(c["km"], 2), passo_da_secondi(passo),
                          durata_leggibile(c["minuti"]), c["data"]])
        else:
            righe.append([etichetta, "—", "nessuna corsa a questa distanza", "—", "—"])

    riga = scrivi_tabella(f, 2, ["Distanza", "Km effettivi", "Passo migliore",
                                 "Durata", "Data"], righe)

    riga += 1
    stile_titolo(f.cell(row=riga, column=1, value="ALTRE STATISTICHE"), 12)

    veloce = min(corse, key=lambda c: c["minuti"] / c["km"])
    lunga  = max(corse, key=lambda c: c["km"])
    salita = max(corse, key=lambda c: c["dislivello"])

    altre = [
        ["Passo più veloce in assoluto",
         passo_da_secondi(veloce["minuti"] * 60 / veloce["km"]) + " min/km",
         f"{veloce['km']} km il {veloce['data']}"],
        ["Uscita più lunga", f"{lunga['km']} km", f"il {lunga['data']}"],
        ["Più dislivello", f"{salita['dislivello']} m", f"il {salita['data']}"],
        ["Totale km percorsi", f"{sum(c['km'] for c in corse):.1f} km", ""],
        ["Totale uscite", len(corse), ""],
    ]
    for i, valori in enumerate(altre):
        r = riga + 1 + i
        f.cell(row=r, column=1, value=valori[0]).font = Font(bold=True)
        f.cell(row=r, column=2, value=valori[1])
        f.cell(row=r, column=3, value=valori[2]).font = Font(color="666666")

# ============================================
# PROGRAMMA PRINCIPALE
# ============================================

if __name__ == "__main__":
    print("\n" + "=" * 50)
    print("   STRAVA — REPORT EXCEL")
    print("=" * 50)

    import sys
    corse = carica_corse(FILE_CSV, dal=None if "--tutto" in sys.argv else DAL)
    mesi = riepilogo_mensile(corse)

    wb = Workbook()
    foglio_dashboard(wb, corse, mesi)
    foglio_corse(wb, corse)
    foglio_mensile(wb, corse, mesi)
    foglio_record(wb, corse)
    wb.save(FILE_XLSX)

    print(f"✓ Report salvato in '{FILE_XLSX}'")
    print(f"  {len(corse)} corse  •  {sum(c['km'] for c in corse):.1f} km  "
          f"•  {len(mesi)} mesi\n")
